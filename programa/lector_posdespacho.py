"""
Lee un posdespacho diario de la AMM (.xlsx) y devuelve:
  - eventos del RSO (texto del Word embebido en RSO1/RSO2/RSO3), con tipo, kV, horas y duración
  - registros de la hoja Indisponibilidades
Solo usa la librería estándar + openpyxl.
"""
import re, zipfile, io, datetime as dt

_SEC = [("ASPECTOS OPERATIVOS", "Operación"), ("SISTEMA EL", "Regional"), ("REGIONAL", "Regional"),
        ("TRANSMISI", "Transmisión"), ("DATOS HIDROL", None)]
_I = re.I
_TIPO = [
    ("Intercambios / NCC-10", re.compile(r"NCC-10|Medidas Administrativas", _I)),
    ("Reserva Rápida (RRa)", re.compile(r"Reserva R[aá]pida|\bR\.?\s?RA\b|\bRRa\b", _I)),
    ("Altas / bajas de centrales", re.compile(r"habilitad\w+ en el Mercado|de baja (del|en el) Mercado", _I)),
    ("Nota regulatoria", re.compile(r"^(Se informa que )?A partir del \d+ de \w+ de \d{4}", _I)),
    ("Redespacho", re.compile(r"cobr\w+ vigencia|redespacho|programa de despacho", _I)),
    ("Generación forzada", re.compile(r"gener\w+ forzad|forzad\w+ (por|para) regulaci", _I)),
    ("Pruebas", re.compile(r"prueba", _I)),
    ("Disparo / falla", re.compile(r"dispar|abiert\w+ por disparo|falla|sin tensi[oó]n|EDAC|bajo voltaje", _I)),
    ("Mantenimiento / maniobra", re.compile(r"desenergizad|mantenimiento|cerrad\w+|abiert\w+|bloquead|deshabilitad|energizad|trabajos|desconectad", _I)),
    ("Disponibilidad de unidades", re.compile(r"(in)?disponib", _I)),
]
_KV = re.compile(r"(\d{2,3}(?:\.\d)?)\s*kV", _I)
_HH = re.compile(r"(?:de|desde las)\s+(\d{1,2}:\d{2})(?::\d{2})?\s*(?:a|hasta las)\s+(\d{1,2}:\d{2})", _I)
_A_LAS = re.compile(r"a las (\d{1,2}:\d{2})", _I)
_CARGA = re.compile(r"sin tensi[oó]n|sin servicio|carga (afectada|desconectada|qued)|desconect\w+ (la )?carga|MW de carga", _I)


def _min(h):
    a, b = h.split(":"); return int(a) * 60 + int(b)


def _paras_docx(blob):
    x = zipfile.ZipFile(io.BytesIO(blob)).read("word/document.xml").decode("utf8", "ignore")
    x = re.sub(r"<w:tbl>.*?</w:tbl>", "", x, flags=re.S)          # tablas (hidrología, cuadros RRa)
    ps = []
    for p in re.findall(r"<w:p[ >].*?</w:p>", x, flags=re.S):
        t = re.sub(r"<[^>]+>", "", re.sub(r"<w:tab/>", " ", p))
        t = re.sub(r"\s+", " ", t.replace("&amp;", "&").replace("&lt;", "<").replace("&gt;", ">")).strip()
        if t: ps.append(t)
    return ps


def _tipo(texto, ambito):
    if ambito == "Regional" and re.search(r"interconexi", texto, _I) and not re.search(r"Reserva R|RRA", texto):
        return "Interconexión"
    if re.search(r"mantenimiento|AMM-[A-Z]{3}\d{2}-", texto, _I) and not re.search(r"dispar|falla", texto, _I):
        return "Mantenimiento / maniobra"
    for nom, rx in _TIPO:
        if rx.search(texto): return nom
    return "Otro"


def leer_rso(path):
    z = zipfile.ZipFile(path)
    # orden de hojas -> archivos sheetN y sus embeddings
    wb = z.read("xl/workbook.xml").decode()
    rels = z.read("xl/_rels/workbook.xml.rels").decode()
    rid2t = dict(re.findall(r'Id="(rId\d+)"[^>]*Target="([^"]+)"', rels)) | {a: b for b, a in re.findall(r'Target="([^"]+)"[^>]*Id="(rId\d+)"', rels)}
    hojas = re.findall(r'<sheet [^>]*name="([^"]+)"[^>]*r:id="(rId\d+)"', wb)
    blobs = []
    for nom, rid in hojas:
        if not nom.upper().startswith("RSO"): continue
        sh = rid2t[rid].split("/")[-1]
        try: r = z.read(f"xl/worksheets/_rels/{sh}.rels").decode()
        except KeyError: continue
        for emb in re.findall(r'Target="\.\./embeddings/([^"]+\.docx)"', r):
            blobs.append(z.read("xl/embeddings/" + emb))
    eventos, amb, pend = [], "Operación", None
    for b in blobs:
        for t in _paras_docx(b):
            T = t.upper()
            hit = next((s for k, s in _SEC if T.startswith(("INFORMACIÓN", "INFORMACION", "ASPECTOS", "DATOS")) and k in T), "x")
            if hit != "x":
                amb = hit; pend = None; continue
            if amb is None or len(t) < 25 or re.match(r"^(CONTIN|REGRESAR|P[aá]gina)", t, _I):
                continue
            if pend is not None and (t.startswith(("·", "-", "•")) or (len(t) < 90 and not re.search(r"\d{1,2}:\d{2}", t[:40]) and eventos)):
                eventos[pend]["texto"] += " · " + t.lstrip("·-• "); continue
            eventos.append({"ambito": amb, "texto": t})
            pend = len(eventos) - 1 if t.rstrip().endswith(":") else None
    for e in eventos:
        t = e["texto"]; e["tipo"] = _tipo(t, e["ambito"])
        kvs = [float(k) for k in _KV.findall(t)]; e["kv"] = max(kvs) if kvs else None
        m = _HH.search(t)
        if m:
            e["h_ini"], e["h_fin"] = m.group(1).zfill(5), m.group(2).zfill(5)
            d = _min(e["h_fin"]) - _min(e["h_ini"]); e["dur_min"] = d if d >= 0 else None
        else:
            m = _A_LAS.search(t); e["h_ini"] = m.group(1).zfill(5) if m else None; e["h_fin"] = None; e["dur_min"] = None
        e["carga_afectada"] = bool(_CARGA.search(t))
        e["mw_gen_perdida"] = None
    return eventos


_CAT = [("No atribuible", "No atribuible"), ("Degrad", "Degradación"), ("Disparo de Unidad", "Disparo de Unidad"),
        ("Falla al ser", "Falla al ser Convocado"), ("Causa Externa", "Causa Externa"), ("Mant", "Mantenimiento"),
        ("Excede", "Mantenimiento"), ("Forzad", "Forzada")]


def leer_indisp(path, fecha):
    import openpyxl
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = next((wb[s] for s in wb.sheetnames if s.lower().startswith("indisp")), None)
    out = []
    if ws is None: return out
    for r in ws.iter_rows(min_row=4, values_only=True):
        if not r or not isinstance(r[0], str) or not re.match(r"^[A-Z0-9&]{2,6}-[A-Z]{1,2}\d*$", r[0].strip()): continue
        u, mot = r[0].strip(), str(r[1] or "")
        sale_f, sale_h, ent_f, ent_h = r[2], r[3], r[4], r[5]
        obs = r[6] if len(r) > 6 else None
        if not isinstance(ent_f, (dt.date, dt.datetime)):   # sin fecha de entrada: columnas corridas
            obs, ent_f, ent_h = ent_f, None, None
        sf = sale_f.date() if isinstance(sale_f, dt.datetime) else sale_f
        cat = next((c for k, c in _CAT if k.lower() in mot.lower()), mot)
        out.append({"unidad": u, "central": re.sub(r"\d+$", "", u), "cat": cat, "tipo": mot,
                    "sale": f"{sf} {sale_h}" if sf else None, "entra": f"{ent_f} {ent_h}" if ent_f else None,
                    "nueva": sf == fecha, "obs": str(obs) if obs else None})
    return out

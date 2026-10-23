"""Lectura del PosDespacho diario de la AMM (archivo dentro de PDyyyymmdd.zip) y comparación con lo programado."""
import re, io, zipfile, datetime as dt
from pathlib import Path
import openpyxl
from parse_pds import (norm, CVG_OVERRIDE, cvg_index, merge_versions, parse_file, MIX_ORDER)

MESES_ES = ['enero', 'febrero', 'marzo', 'abril', 'mayo', 'junio', 'julio', 'agosto',
            'septiembre', 'octubre', 'noviembre', 'diciembre']

# ---------------------------------------------------------------- posdespacho
def _f(x):
    try: return float(x)
    except (TypeError, ValueError): return None


def _docx_text_tables(blob):
    """Devuelve las filas de las tablas de un .docx embebido (lista de listas de texto)."""
    z = zipfile.ZipFile(io.BytesIO(blob))
    xml = z.read('word/document.xml').decode('utf8')
    rows = []
    for tr in re.findall(r'<w:tr[ >].*?</w:tr>', xml, re.S):
        cells = [re.sub(r'<[^>]+>', '', tc).strip() for tc in re.findall(r'<w:tc[ >].*?</w:tc>', tr, re.S)]
        rows.append(cells)
    return rows


def _cota_chixoy(xlsx_path):
    """La cota viene en un documento de Word incrustado en la hoja RSO1 (tabla DATOS HIDROLÓGICOS)."""
    with zipfile.ZipFile(xlsx_path) as z:
        for name in z.namelist():
            if name.startswith('xl/embeddings/') and name.endswith('.docx'):
                for row in _docx_text_tables(z.read(name)):
                    if row and norm(row[0]).startswith('CHIXOY'):
                        vals = [_f(c.replace(',', '')) for c in row[1:]]
                        return {'cota': vals[0] if vals else None,
                                'q_ent': vals[1] if len(vals) > 1 else None,
                                'q_turb': vals[2] if len(vals) > 2 else None}
    return {}


def parse_pos(path):
    path = Path(path)
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    out = {'file': path.name}

    # Fecha (del nombre PDyyyymmdd o del texto de la hoja POE)
    m = re.search(r'(\d{8})', path.name)
    if m:
        out['date'] = dt.datetime.strptime(m.group(1), '%Y%m%d').date()
    else:
        txt = ' '.join(str(c) for r in wb['POE'].iter_rows(max_row=5, values_only=True) for c in r if c)
        mm = re.search(r'(\d{1,2}) de (\w+) del? (\d{4})', txt.lower())
        out['date'] = dt.date(int(mm.group(3)), MESES_ES.index(mm.group(2)) + 1, int(mm.group(1)))

    # Carga Horaria: fila con nemos y fila siguiente con la energía del día (MWh)
    rows = list(wb['Carga Horaria'].iter_rows(max_row=12, values_only=True))
    hdr = next(i for i, r in enumerate(rows) if any(str(c).strip() == 'Total Generado' for c in r if c))
    real = {}
    for name, val in zip(rows[hdr], rows[hdr + 1]):
        if name is not None and str(name).strip():
            real[str(name).strip()] = _f(val) or 0.0
    out['total_gen'] = real.pop('Total Generado', None)
    out['dem_real'] = real.pop('Demanda SNI', None)
    out['real'] = real

    # POE horario
    poe = []
    for r in wb['POE'].iter_rows(values_only=True):
        r = [c for c in r if c is not None]
        if len(r) >= 4 and isinstance(_f(r[3]), float) and isinstance(r[2], str) and ':' in str(r[0]):
            poe.append({'de': str(r[0])[:5], 'gen': r[2].strip(), 'poe': _f(r[3])})
    out['poe'] = poe

    # Demanda máxima (Curva S.N.I.)
    out['dmax'] = out['dmax_hora'] = None
    for r in wb['Curva S.N.I.'].iter_rows(max_row=12, values_only=True):
        vals = [c for c in r if c is not None]
        if len(vals) >= 2 and _f(vals[0]) and _f(vals[0]) > 500 and _f(vals[1]) is not None:
            out['dmax'] = _f(vals[0])
            h = _f(vals[1]); hh, mm_ = int(h), int(round((h - int(h)) * 100))
            out['dmax_hora'] = f'{hh:02d}:{mm_:02d}'
            break

    # Caudales de Chixoy (RSO3)
    chx = {}
    rows = list(wb['RSO3'].iter_rows(values_only=True))
    for i, r in enumerate(rows):
        cells = [str(c).strip() if c is not None else '' for c in r]
        if 'CHX-H1' in cells or 'CHX-H' in cells:
            col = cells.index('CHX-H1') if 'CHX-H1' in cells else cells.index('CHX-H')
            for r2 in rows[i + 1:]:
                lab = norm(next((c for c in r2 if c is not None), ''))
                v = _f(r2[col]) if col < len(r2) else None
                if lab.startswith('Q PROMEDIO'): chx['q_prom'] = v
                elif lab.startswith('CAUDAL MAX'): chx['q_max'] = v
                elif lab.startswith('CAUDAL MIN'): chx['q_min'] = v
            break
    chx.update({k: v for k, v in _cota_chixoy(path).items() if v is not None and k not in chx})
    out['chixoy'] = chx
    return out


def extract_zip(zip_bytes, dest_dir, date):
    """Guarda el Excel del zip como PDyyyymmdd.xlsx."""
    z = zipfile.ZipFile(io.BytesIO(zip_bytes))
    x = [n for n in z.namelist() if n.lower().endswith(('.xlsx', '.xlsm'))]
    if not x:
        raise ValueError('El zip no trae Excel')
    dest = Path(dest_dir) / f'PD{date:%Y%m%d}.xlsx'
    dest.write_bytes(z.read(x[0]))
    return dest


# ---------------------------------------------------------------- comparación
EXTRA_NEMO = {('MOTORES', 'INNOVA'): 'INV-G', ('EOLICA', 'COMAPA'): 'CMP-E', ('MOTORES', 'ESUS'): 'ESU-G'}
# Plantas que se suman como una sola (cascadas, bloques, unidades). Lo demás se agrupa por
# nemo sin número de unidad/bloque (p. ej. JEN-C1 + JEN-C2 = JEN-C; MAG-B1..B7 = MAG-B).
GRUPOS = {
    'RENACE': ['REN-H', 'RE2-H', 'RE3-H', 'RE4-H'],
    'INDE': ['CHX-H', 'AGU-H', 'JUR-H', 'LES-H', 'SMA-H', 'POR-H', 'PAL-H'],
    'XACBAL': ['XAC-H', 'XAD-H'],
}
NOMBRE_GRUPO = {'JEN-C': 'JAGUAR ENERGY', 'MAG-B': 'MAGDALENA', 'PNT-B': 'PANTALEÓN', 'TND-B': 'TRINIDAD',
                'TUL-B': 'TULULÁ', 'LPA-B': 'LAS PALMAS BUNKER', 'TDL-B': 'TEXTILES DEL LAGO', 'TAM-G': 'TAMPA',
                'MNL-H': 'EL MANANTIAL', 'HRU-F': 'HORUS', 'TER-B': 'TÉRMICA', 'SAA-B': 'SANTA ANA',
                'PGO-B': 'PALO GORDO', 'EPI-B': 'EL PILAR', 'LUN-B': 'LA UNIÓN'}
_G = {k: g for g, ks in GRUPOS.items() for k in ks}
TECH_SUFIJO = {'H': 'HIDRO', 'F': 'SOLAR', 'E': 'EOLICA', 'C': 'VAPOR', 'V': 'VAPOR'}


def gkey(nemo):
    m = re.match(r'([A-Z0-9&]+)-([A-Z])', nemo)
    k = f'{m.group(1)}-{m.group(2)}' if m else nemo
    return _G.get(k, k)


def base(nemo):
    return re.sub(r'\d+$', '', nemo)


def comparar(pos, prog):
    """pos = parse_pos(); prog = semana vigente (merge_versions). Devuelve el resumen del día."""
    d = pos['date']
    i = prog['days'].index(d)
    _, by_name = cvg_index(prog['cvg'])

    # Programado por planta y mapa nemo -> tecnología / nombre
    tech_of, name_of, P_plant = {}, {}, {}
    P_tech = {}
    for p in prog['plants']:
        if p['tech'].startswith('MER:'): continue
        v = p['daily'][i]
        P_tech[p['tech']] = P_tech.get(p['tech'], 0) + v
        nemo = EXTRA_NEMO.get((p['section'], p['n'])) or CVG_OVERRIDE.get((p['section'], p['n'])) or by_name.get(p['n'])
        if not nemo: continue
        tech_of.setdefault(base(nemo), p['tech']); tech_of.setdefault(nemo, p['tech'])
        g = gkey(nemo)
        P_plant[g] = P_plant.get(g, 0) + v
        name_of.setdefault(g, NOMBRE_GRUPO.get(g, g if g in GRUPOS else p['name']))

    # Real por tecnología y por planta
    R_tech, R_plant, sin_tech = {}, {}, []
    real = dict(pos['real'])
    mex = -real.pop('MEX-I', 0.0)         # importación desde México viene negativa
    real.pop('SER-I', None)               # intercambio MER: no se compara (igual que el mix del predespacho)
    for nemo, v in real.items():
        t = tech_of.get(nemo) or tech_of.get(base(nemo))
        if not t:
            m = re.match(r'[A-Z0-9&]+-([A-Z])', nemo)
            t = TECH_SUFIJO.get(m.group(1)) if m else None
            if nemo.startswith('CHX'): t = 'CHIXOY'
        if not t:
            if abs(v) > 0.5: sin_tech.append(nemo)
            continue
        R_tech[t] = R_tech.get(t, 0) + v
        g = gkey(nemo)
        R_plant[g] = R_plant.get(g, 0) + v
        name_of.setdefault(g, NOMBRE_GRUPO.get(g, nemo))
    # México: Energía del Caribe (contrato fijo) + importación de oportunidad
    caribe_prog = P_tech.get('CARIBE', 0.0)
    R_tech['CARIBE'] = min(max(mex, 0.0), caribe_prog)
    R_tech['MEX'] = max(mex - caribe_prog, 0.0)
    P_tech.setdefault('MEX', 0.0)

    techs = [t for t in MIX_ORDER if R_tech.get(t, 0) > 0.5 or P_tech.get(t, 0) > 0.5]
    tech_rows = [{'tech': t, 'real': R_tech.get(t, 0.0), 'prog': P_tech.get(t, 0.0)} for t in techs]
    for r in tech_rows: r['diff'] = r['real'] - r['prog']
    no_chx = [r for r in tech_rows if r['tech'] != 'CHIXOY']
    pos3 = sorted([r for r in no_chx if r['diff'] > 0.5], key=lambda r: -r['diff'])[:3]
    neg3 = sorted([r for r in no_chx if r['diff'] < -0.5], key=lambda r: r['diff'])[:3]
    chx_row = next((r for r in tech_rows if r['tech'] == 'CHIXOY'), None)

    plants = []
    for g in set(R_plant) | set(P_plant):
        if g == 'EDC-I': continue
        rv, pv = R_plant.get(g, 0.0), P_plant.get(g, 0.0)
        plants.append({'g': g, 'name': name_of.get(g, g), 'real': rv, 'prog': pv, 'diff': rv - pv})
    top_plants = sorted(plants, key=lambda r: -abs(r['diff']))[:5]

    poe = pos['poe']
    pvals = [x['poe'] for x in poe]
    pmax = max(poe, key=lambda x: x['poe']) if poe else None
    pmin = min(poe, key=lambda x: x['poe']) if poe else None
    return {
        'date': d, 'week': prog['week'], 'prog_file': prog['file'],
        'total_gen': pos['total_gen'], 'dem_real': pos['dem_real'], 'dem_prog': prog['dem_energy'][i],
        'dmax': pos['dmax'], 'dmax_hora': pos['dmax_hora'],
        'poe': poe, 'poe_prom': sum(pvals) / len(pvals) if pvals else None, 'poe_max': pmax, 'poe_min': pmin,
        'chixoy': pos['chixoy'], 'chx_row': chx_row,
        'tech_rows': tech_rows, 'pos3': pos3, 'neg3': neg3, 'top_plants': top_plants, 'sin_tech': sin_tech,
    }
